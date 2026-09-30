"""File-type detection and plain-text extraction for uploaded documents.

Detection uses the file's bytes, never its name or the browser-supplied type.
"""

import io
import re
import zipfile

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .models import CandidateDocument

# DOCX is a zip; cap the uncompressed size to reject zip bombs before parsing.
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
# Below this many characters a PDF is almost certainly scanned images.
MIN_TEXT_CHARS = 200


class DocumentRejected(Exception):
    """The file can't be used. `code` is stable; `message` is safe to show users."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def detect_content_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return CandidateDocument.ContentType.PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" in names:
                    total = sum(info.file_size for info in archive.infolist())
                    if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                        raise DocumentRejected("docx_too_large", "This Word document is too large to process.")
                    return CandidateDocument.ContentType.DOCX
        except zipfile.BadZipFile:
            pass
    raise DocumentRejected("unsupported_type", "Upload a PDF or Word (.docx) file.")


def extract_text(data: bytes, content_type: str, *, max_pdf_pages: int) -> str:
    if content_type == CandidateDocument.ContentType.PDF:
        text = _extract_pdf(data, max_pdf_pages=max_pdf_pages)
    elif content_type == CandidateDocument.ContentType.DOCX:
        text = _extract_docx(data)
    else:
        raise DocumentRejected("unsupported_type", "Upload a PDF or Word (.docx) file.")

    text = _normalise(text)
    if len(text) < MIN_TEXT_CHARS:
        if content_type == CandidateDocument.ContentType.PDF:
            raise DocumentRejected(
                "no_text",
                "We couldn't read any text in this PDF. It may be a scanned image; "
                "please upload a PDF with selectable text or a Word document.",
            )
        raise DocumentRejected("no_text", "This document doesn't contain enough text to read.")
    return text


def _extract_pdf(data: bytes, *, max_pdf_pages: int) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentRejected("encrypted", "This PDF is password-protected. Please upload an unprotected copy.")
        if len(reader.pages) > max_pdf_pages:
            raise DocumentRejected("too_many_pages", f"Resumes can be at most {max_pdf_pages} pages.")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentRejected:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentRejected("unreadable", "This PDF couldn't be read. It may be damaged.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of parser errors
        raise DocumentRejected("unreadable", "This Word document couldn't be read. It may be damaged.") from exc

    parts = [p.text for p in document.paragraphs]
    # Many resumes lay out sections in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _normalise(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
