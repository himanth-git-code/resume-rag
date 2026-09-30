"""Builders for small, real PDF/DOCX files used in tests."""

import io

import docx
from pypdf import PdfWriter

RESUME_LINES = [
    "Jane Doe - Senior Backend Engineer",
    "jane@example.com | Berlin, Germany",
    "Experience: Acme Corp, Senior Backend Engineer, Jan 2020 - Present.",
    "Built Django and PostgreSQL services handling payments for 2 million users.",
    "Introduced Redis caching and cut p95 API latency by 40 percent.",
    "Education: BSc Computer Science, Technical University of Munich, 2016.",
    "Skills: Python, Django, PostgreSQL, Redis, AWS, Docker, Kubernetes.",
]


def text_pdf(lines=RESUME_LINES) -> bytes:
    """A one-page PDF whose text pypdf can extract."""

    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    body = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({esc(line)}) Tj T*" for line in lines) + " ET"
    stream = body.encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def blank_pdf(pages=1) -> bytes:
    """A PDF with no text layer, like a scanned image."""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=842)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def text_docx(lines=RESUME_LINES) -> bytes:
    document = docx.Document()
    for line in lines[:-1]:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = lines[-1]
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
